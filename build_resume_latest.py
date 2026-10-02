"""Generate a JD-tailored resume DOCX and report ATS-style keyword match scores.

Usage examples:
  python3.14 build_resume_latest.py --jd-file jd_input.txt
  python3.14 build_resume_latest.py --jd-text "...paste JD here..."
  cat jd_input.txt | python3.14 build_resume_latest.py
"""
from __future__ import annotations

import argparse
import html
import re
import sys
from collections import OrderedDict
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib.enums import TA_CENTER  # pyright: ignore[reportMissingImports]
from reportlab.lib.pagesizes import A4  # pyright: ignore[reportMissingImports]
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet  # pyright: ignore[reportMissingImports]
from reportlab.lib.units import cm  # pyright: ignore[reportMissingImports]
from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate  # pyright: ignore[reportMissingImports]

SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_JD_FILE = SCRIPT_DIR / "jd_input.txt"
DEFAULT_OUTPUT = SCRIPT_DIR / "Sai_Charan_Tumpuri_Resume_latest.docx"
DEFAULT_PDF_OUTPUT = DEFAULT_OUTPUT.with_suffix(".pdf")

STOPWORDS = {
    "a",
    "an",
    "and",
    "are",
    "as",
    "at",
    "be",
    "by",
    "for",
    "from",
    "have",
    "has",
    "had",
    "in",
    "into",
    "is",
    "it",
    "its",
    "of",
    "on",
    "or",
    "our",
    "that",
    "the",
    "their",
    "this",
    "to",
    "with",
    "we",
    "you",
    "your",
    "will",
    "work",
    "working",
    "experience",
    "years",
    "year",
    "role",
    "roles",
    "responsibilities",
    "responsibility",
    "team",
    "teams",
    "company",
    "companies",
    "candidate",
    "candidates",
    "cloud",
    "devops",
    "developer",
    "developers",
    "engineer",
    "engineers",
    "engineering",
    "engineered",
    "preferred",
    "required",
    "strong",
    "platform",
    "platforms",
    "solution",
    "solutions",
    "software",
    "system",
    "systems",
    "service",
    "services",
    "application",
    "applications",
    "technical",
    "technology",
    "technologies",
    "senior",
    "junior",
    "principal",
    "lead",
    "manager",
    "management",
    "call",
    "on-call",
    "incident",
    "response",
    "postmortem",
    "postmortems",
    "slo",
    "slos",
    "sla",
    "slas",
    "ability",
    "skills",
    "skill",
    "knowledge",
    "familiar",
    "familiarity",
    "include",
    "including",
    "preferred",
    "ability",
    "hands",
    "hands-on",
    "handson",
    "build",
    "building",
    "manage",
    "managing",
    "maintain",
    "maintaining",
    "support",
    "supporting",
    "develop",
    "developing",
    "design",
    "designing",
    "implement",
    "implementing",
    "create",
    "creating",
    "use",
    "using",
    "knowledge",
    "understanding",
    "ability",
    "proficiency",
    "proficient",
    "excellent",
    "good",
    "must",
    "nice",
    "plus",
    "bonus",
    "want",
    "wanted",
    "looking",
    "seeking",
    "etc",
}

TERM_ALIASES = OrderedDict(
    [
        ("gcp", ["gcp", "google cloud", "google cloud platform"]),
        ("aws", ["aws", "amazon web services"]),
        ("azure", ["azure", "microsoft azure"]),
        ("kubernetes", ["kubernetes", "k8s"]),
        ("terraform", ["terraform"]),
        ("terragrunt", ["terragrunt"]),
        ("terraform modules", ["terraform modules"]),
        ("terragrunt modules", ["terragrunt modules"]),
        ("ansible", ["ansible"]),
        ("packer", ["packer"]),
        ("docker", ["docker"]),
        ("helm", ["helm"]),
        ("kustomize", ["kustomize"]),
        ("gitlab", ["gitlab"]),
        ("gocd", ["gocd"]),
        ("ci cd", ["ci cd", "continuous integration", "continuous delivery", "continuous deployment"]),
        ("gitlab ci cd", ["gitlab ci cd", "gitlab ci/cd"]),
        ("github actions", ["github actions"]),
        ("cloud build", ["cloud build"]),
        ("jenkins", ["jenkins"]),
        ("argocd", ["argocd"]),
        ("prometheus", ["prometheus"]),
        ("grafana", ["grafana"]),
        ("cloud monitoring", ["cloud monitoring", "google cloud monitoring"]),
        ("cloudwatch", ["cloudwatch", "aws cloudwatch"]),
        ("log pipelines", ["log pipelines"]),
        ("open telemetry", ["open telemetry", "opentelemetry"]),
        ("distributed tracing", ["distributed tracing", "tracing"]),
        ("logging", ["logging", "centralized logging", "log aggregation"]),
        ("elk stack", ["elk stack", "elk", "elasticsearch", "logstash", "kibana"]),
        ("on call", ["on call", "on-call"]),
        ("slo", ["slo", "slos", "service level objective", "service level objectives"]),
        ("sla", ["sla", "slas", "service level agreement", "service level agreements"]),
        ("postmortem", ["postmortem", "postmortems", "incident review"]),
        ("incident response", ["incident response", "incident management"]),
        ("observability", ["observability"]),
        ("scalability", ["scalability"]),
        ("capacity planning", ["capacity planning"]),
        ("linux", ["linux"]),
        ("linux-based", ["linux-based", "linux based"]),
        ("networking", ["networking", "tcp ip", "dns", "load balancing", "vpc"]),
        ("mysql", ["mysql"]),
        ("cassandra", ["cassandra"]),
        ("secrets management", ["secrets management"]),
        ("configuration management", ["configuration management"]),
        ("environment setup", ["environment setup"]),
        ("database clusters", ["database clusters", "database cluster", "db clusters"]),
        ("environment-specific access control", ["environment-specific access control"]),
        ("aws best practices", ["aws best practices"]),
        ("ssm", ["ssm", "aws ssm", "systems manager"]),
        ("hybrid cloud", ["hybrid cloud"]),
        ("hybrid cloud environments", ["hybrid cloud environments"]),
        ("high availability", ["high availability", "ha"]),
        ("high-availability environments", ["high-availability environments", "high availability environments"]),
        ("linux system administration", ["linux system administration", "linux administration", "system administration"]),
        ("shell scripting", ["shell scripting", "shell script", "bash scripting"]),
        ("containerization", ["containerization", "containerisation"]),
        ("aws cloudformation", ["aws cloudformation", "cloudformation"]),
        ("ansible roles", ["ansible roles"]),
        ("ansible modules", ["ansible modules"]),
        ("access controls", ["access controls", "access control"]),
        ("encryption", ["encryption"]),
        ("vulnerability management", ["vulnerability management", "vulnerability scans", "vulnerability scanning"]),
        ("regulated", ["regulated"]),
        ("fintech", ["fintech"]),
        ("bigquery", ["bigquery"]),
        ("dataflow", ["dataflow"]),
        ("airflow", ["airflow", "composer"]),
        ("apache beam", ["apache beam"]),
        ("gke", ["gke"]),
        ("eks", ["eks"]),
        ("gcs", ["gcs"]),
        ("iam", ["iam"]),
        ("vpc", ["vpc"]),
        ("microservices", ["microservices"]),
        ("root cause analysis", ["root cause analysis", "root cause"]),
        ("cost optimization", ["cost optimization", "cost optimization"]),
    ]
)

DISPLAY_TERM = {
    "ci cd": "CI/CD",
    "open telemetry": "OpenTelemetry",
    "on call": "on-call",
    "gcp": "GCP",
    "aws": "AWS",
    "azure": "Azure",
    "gke": "GKE",
    "eks": "EKS",
    "gocd": "GoCD",
    "gcs": "GCS",
    "iam": "IAM",
    "vpc": "VPC",
    "slo": "SLO",
    "sla": "SLA",
    "gcs": "GCS",
    "ci cd": "CI/CD",
    "apache beam": "Apache Beam",
    "argocd": "ArgoCD",
    "k8s": "K8s",
}

RESUME_TEMPLATE = {
    "name": "SAI CHARAN TUMPURI",
    "title": "DevOps / Site Reliability Engineer",
    "contact": {
        "prefix": "post2saicharan@gmail.com  |  +91 9390144066  |  Chennai, India  |  ",
        "links": [
            ("GitHub", "https://github.com/saicharancodes"),
            ("LinkedIn", "https://www.linkedin.com/in/sai-charan-57b049232/"),
            ("Portfolio", "https://saicharancodes.github.io/portfolio/"),
        ],
    },
    "summary": (
        "Platform/DevOps engineer with 3+ years designing and automating scalable cloud "
        "infrastructure on GCP and AWS across Linux-based, hybrid environments for 20+ teams "
        "at Sky (Comcast). Hands-on with Linux system administration, shell scripting, Terraform "
        "and Ansible, Kubernetes (GKE/EKS), Docker, and CI/CD across Jenkins, GitHub Actions, "
        "Cloud Build, and ArgoCD. Own on-call incident response, monitoring, logging, and security "
        "lifecycle tasks including access controls and vulnerability management. Built skyform, an "
        "internal IaC self-service platform that cut infra ticket resolution time by 60%. AWS "
        "Solutions Architect - Associate and GCP Professional Cloud Architect certified."
    ),
    "skills": [
        {
            "label": "Cloud",
            "base": "GCP, AWS",
            "pool": ["gcp", "aws"],
        },
        {
            "label": "Systems & Networking",
            "base": "Linux system administration, shell scripting, TCP/IP networking, DNS, load balancing, VPC, hybrid cloud",
            "pool": ["linux system administration", "shell scripting", "linux", "networking", "dns", "load balancing", "vpc", "hybrid cloud"],
        },
        {
            "label": "Containers & Orchestration",
            "base": "Kubernetes (GKE, EKS), Docker, Helm",
            "pool": ["kubernetes", "gke", "eks", "docker", "containerization", "helm", "kustomize"],
        },
        {
            "label": "Infrastructure as Code",
            "base": "Terraform, Ansible, Packer, configuration management, environment setup",
            "pool": ["terraform", "ansible", "ansible roles", "ansible modules", "packer", "configuration management", "environment setup", "aws cloudformation"],
        },
        {
            "label": "CI/CD & GitOps",
            "base": "Jenkins, GitHub Actions, Cloud Build, ArgoCD, GitLab CI/CD",
            "pool": ["jenkins", "github actions", "cloud build", "argocd", "ci cd", "gitlab ci cd"],
        },
        {
            "label": "Observability & Incident Response",
            "base": "Prometheus, Grafana, Cloud Monitoring, OpenTelemetry, distributed tracing, logging, log pipelines, on-call, SLO/SLA, postmortems",
            "pool": [
                "prometheus",
                "grafana",
                "cloud monitoring",
                "cloudwatch",
                "open telemetry",
                "distributed tracing",
                "logging",
                "elk stack",
                "log pipelines",
                "on call",
                "slo",
                "sla",
                "postmortem",
                "incident response",
                "observability",
                "vulnerability management",
            ],
        },
        {
            "label": "Languages",
            "base": "Python, Go, Bash, Groovy, SQL",
            "pool": ["python", "go", "bash", "groovy", "sql"],
        },
        {
            "label": "Data Platform",
            "base": "BigQuery, Dataflow (Apache Beam), Airflow/Composer",
            "pool": ["bigquery", "dataflow", "apache beam", "airflow", "composer"],
        },
        {
            "label": "Security",
            "base": "IAM, access controls, encryption, vulnerability management",
            "pool": ["iam", "access controls", "encryption", "vulnerability management"],
        },
    ],
    "experience": [
        {
            "role": "DevOps Engineer II",
            "company": "Comcast",
            "dates": "Mar 2024 - Present",
            "location": "Chennai, India",
            "bullets": [
                "Designed and operated Kubernetes-native CI/CD on GKE with ephemeral pod agents (Groovy + Python DSL) across 20+ pipelines, cutting idle compute by 30% and integrating E2E tests, artifact promotion, and vulnerability scans.",
                "Built skyform, an internal Terraform abstraction with project-isolated remote state that lets engineers self-serve GCP infra (BigQuery, GKE, Dataflow); reduced infra ticket resolution time by 60% and unblocked 20+ data pipelines.",
                "Migrated legacy Dataflow and batch workloads to GKE using Helm and HPA, standardized chart templates, and drove capacity planning, scalability, and high availability improvements by enforcing resource requests/limits - reducing job runtime by 25%.",
                "Owned OS and security lifecycle: rebuilt GCP golden images with Packer, led migration of 300+ VMs from CentOS 7 to CentOS 9 with zero SLA breaches, and decommissioned 40+ underutilized VMs to cut cost and CVE exposure.",
            ],
        },
        {
            "role": "DevOps Engineer I",
            "company": "Comcast",
            "dates": "Jun 2023 - Mar 2024",
            "location": "Chennai, India",
            "bullets": [
                "Provisioned secure GCP infrastructure (VMs, IAM, BigQuery, GCS, VPC networking) using Terraform and Ansible with zero-touch Cloud Build CI/CD, supporting Linux-based environment setup and shell-scripted automation while cutting deployment time by 60%.",
                "Integrated InfraCost into GitHub PR checks for automated cost visibility on IaC changes, enabling proactive budget forecasting.",
                "Rebuilt 30+ legacy projects and decommissioned obsolete VMs; authored runbooks for configuration management, environment setup, and incident response.",
                "Owned incident management for 300+ production Linux VMs on on-call rotation, monitoring CPU, memory, and 5xx alerts via Cloud Monitoring and Grafana; resolved P1-P4 incidents against SLOs and authored postmortems to drive toil reduction.",
            ],
        },
    ],
    "certifications": [
        (
            "AWS Certified Solutions Architect - Associate (Sep 2025)",
            "https://www.credly.com/badges/1ef62e52-3d2c-4052-9254-19ea8275f0c1/public_url",
        ),
        (
            "Google Cloud Professional Cloud Architect (Feb 2026)",
            "https://www.credly.com/badges/0da7271f-b69b-4cfb-bc1c-9978cae5820c/public_url",
        ),
    ],
    "education": {
        "school": "SASTRA University - B.Tech, Electronics & Communication Engineering",
        "dates": "Jul 2019 - Jun 2023  |  GPA: 7.93",
    },
}


def normalize_text(text: str) -> str:
    text = text.lower().replace("/", " ")
    text = re.sub(r"[^a-z0-9+#\-\s]", " ", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def display_term(term: str) -> str:
    key = normalize_text(term)
    if key in DISPLAY_TERM:
        return DISPLAY_TERM[key]
    return term.replace("  ", " ").strip().title()


def term_variants(term: str) -> list[str]:
    key = normalize_text(term)
    variants = [key]
    if key in TERM_ALIASES:
        variants.extend(TERM_ALIASES[key])
    return sorted(set(normalize_text(variant) for variant in variants if variant))


def contains_term(term: str, text: str) -> bool:
    normalized_text = normalize_text(text)
    for variant in term_variants(term):
        if variant and variant in normalized_text:
            return True
    return False


def extract_jd_terms(jd_text: str) -> set[str]:
    normalized = normalize_text(jd_text)
    found: set[str] = set()

    for canonical, aliases in TERM_ALIASES.items():
        if any(alias in normalized for alias in aliases):
            found.add(canonical)

    return found


def score_text(text: str, jd_terms: set[str]) -> tuple[float, set[str]]:
    hits: set[str] = set()
    score = 0.0
    for term in jd_terms:
        if contains_term(term, text):
            hits.add(term)
            score += 1.0
            if " " in term:
                score += 0.25
    return score, hits


def build_skill_value(base_value: str, pool: list[str], jd_terms: set[str]) -> str:
    value = base_value
    seen = normalize_text(base_value)
    additions: list[str] = []
    for term in pool:
        canonical = normalize_text(term)
        if canonical in jd_terms and canonical not in seen:
            additions.append(display_term(canonical))
            seen += " " + canonical
    if additions:
        value = f"{base_value}, " + ", ".join(additions)
    return value


def rank_bullets(bullets: list[str], jd_terms: set[str]) -> list[str]:
    scored = []
    for index, bullet in enumerate(bullets):
        score, hits = score_text(bullet, jd_terms)
        scored.append((score, len(hits), index, bullet))
    scored.sort(key=lambda item: (-item[0], -item[1], item[2]))
    return [item[3] for item in scored]


def build_tailored_resume(jd_terms: set[str]) -> dict:
    tailored = deepcopy(RESUME_TEMPLATE)
    for skill in tailored["skills"]:
        skill["display"] = build_skill_value(skill["base"], skill["pool"], jd_terms)
    for role in tailored["experience"]:
        role["bullets"] = rank_bullets(role["bullets"], jd_terms)
    return tailored


def add_horizontal_line(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "666666")
    pbdr.append(bottom)
    p_pr.append(pbdr)


def add_section_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)
    add_horizontal_line(p)
    return p


def add_skill_line(doc, label, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    r1 = p.add_run(f"{label}: ")
    r1.bold = True
    r1.font.size = Pt(10)
    r2 = p.add_run(value)
    r2.font.size = Pt(10)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Cm(0.5)
    run = p.runs[0] if p.runs else p.add_run()
    if not p.runs:
        run = p.add_run(text)
    else:
        run.text = text
    run.font.size = Pt(10)


def add_hyperlink(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1155CC")
    rPr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "20")
    rPr.append(sz)
    new_run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_bullet_with_link(doc, text, url, link_text="Badge"):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Cm(0.5)
    r = p.add_run(text + "  |  ")
    r.font.size = Pt(10)
    add_hyperlink(p, url, link_text)


def add_role_header(doc, role, company, dates, location):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(17.5), WD_ALIGN_PARAGRAPH.RIGHT)
    r = p.add_run(f"{role} - {company}")
    r.bold = True
    r.font.size = Pt(10.5)
    r2 = p.add_run(f"\t{dates}")
    r2.bold = True
    r2.font.size = Pt(10)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(1)
    r3 = p2.add_run(location)
    r3.italic = True
    r3.font.size = Pt(9.5)


def build_document(profile: dict, output_path: Path) -> None:
    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(0.8)
        section.bottom_margin = Cm(0.8)
        section.left_margin = Cm(1.4)
        section.right_margin = Cm(1.4)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)

    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.paragraph_format.space_after = Pt(0)
    nr = name.add_run(profile["name"])
    nr.bold = True
    nr.font.size = Pt(18)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(0)
    tr = title.add_run(profile["title"])
    tr.font.size = Pt(11)

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(2)
    cr = contact.add_run(profile["contact"]["prefix"])
    cr.font.size = Pt(9.5)
    for index, (label, url) in enumerate(profile["contact"]["links"]):
        add_hyperlink(contact, url, label)
        if index < len(profile["contact"]["links"]) - 1:
            spacer = contact.add_run("  |  ")
            spacer.font.size = Pt(9.5)

    add_section_heading(doc, "Summary")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    p.add_run(profile["summary"]).font.size = Pt(9.5)

    add_section_heading(doc, "Technical Skills")
    for skill in profile["skills"]:
        add_skill_line(doc, skill["label"], skill["display"])

    add_section_heading(doc, "Professional Experience")
    for role in profile["experience"]:
        add_role_header(doc, role["role"], role["company"], role["dates"], role["location"])
        for bullet in role["bullets"]:
            add_bullet(doc, bullet)

    add_section_heading(doc, "Certifications")
    for cert_text, cert_url in profile["certifications"]:
        add_bullet_with_link(doc, cert_text, cert_url)

    add_section_heading(doc, "Education")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(profile["education"]["school"])
    r.bold = True
    r.font.size = Pt(10)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(profile["education"]["dates"])
    r2.font.size = Pt(9.5)
    r2.italic = True

    doc.save(output_path)


def build_pdf_document(profile: dict, output_path: Path) -> None:
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ResumeTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=20,
        alignment=TA_CENTER,
        spaceAfter=2,
    )
    subtitle_style = ParagraphStyle(
        "ResumeSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=11,
        leading=13,
        alignment=TA_CENTER,
        spaceAfter=2,
    )
    contact_style = ParagraphStyle(
        "ResumeContact",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=11,
        alignment=TA_CENTER,
        spaceAfter=6,
    )
    section_style = ParagraphStyle(
        "SectionHeading",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=12,
        textColor="#1F2A44",
        spaceBefore=4,
        spaceAfter=2,
    )
    body_style = ParagraphStyle(
        "Body",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9.2,
        leading=11,
        spaceAfter=2,
    )
    role_style = ParagraphStyle(
        "Role",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=10,
        leading=11,
        spaceBefore=2,
        spaceAfter=0,
    )
    location_style = ParagraphStyle(
        "Location",
        parent=styles["Normal"],
        fontName="Helvetica-Oblique",
        fontSize=8.5,
        leading=10,
        spaceAfter=1,
    )
    bullet_style = ParagraphStyle(
        "Bullet",
        parent=body_style,
        leftIndent=10,
        firstLineIndent=0,
        bulletIndent=0,
        spaceAfter=1,
    )

    story = [
        Paragraph(html.escape(profile["name"]), title_style),
        Paragraph(html.escape(profile["title"]), subtitle_style),
    ]

    contact_text = html.escape(profile["contact"]["prefix"]) + ""
    contact_text += " | ".join(
        f'<font color="blue"><link href="{html.escape(url)}"><u>{html.escape(label)}</u></link></font>'
        for label, url in profile["contact"]["links"]
    )
    story.append(Paragraph(contact_text, contact_style))

    def add_section(title: str) -> None:
        story.append(Paragraph(title.upper(), section_style))
        story.append(HRFlowable(width="100%", thickness=0.8, color="#666666", spaceBefore=0, spaceAfter=3))

    def add_bullet_paragraph(text: str) -> None:
        story.append(Paragraph(f"• {html.escape(text)}", bullet_style))

    add_section("Summary")
    story.append(Paragraph(html.escape(profile["summary"]), body_style))

    add_section("Technical Skills")
    for skill in profile["skills"]:
        story.append(Paragraph(f"<b>{html.escape(skill['label'])}:</b> {html.escape(skill['display'])}", body_style))

    add_section("Professional Experience")
    for role in profile["experience"]:
        story.append(Paragraph(f"{html.escape(role['role'])} - {html.escape(role['company'])}\t{html.escape(role['dates'])}", role_style))
        story.append(Paragraph(html.escape(role["location"]), location_style))
        for bullet in role["bullets"]:
            add_bullet_paragraph(bullet)

    add_section("Certifications")
    for cert_text, cert_url in profile["certifications"]:
        story.append(Paragraph(f'• {html.escape(cert_text)} | <font color="blue"><link href="{html.escape(cert_url)}"><u>Badge</u></link></font>', bullet_style))

    add_section("Education")
    story.append(Paragraph(f"<b>{html.escape(profile['education']['school'])}</b>", body_style))
    story.append(Paragraph(html.escape(profile["education"]["dates"]), body_style))

    doc = SimpleDocTemplate(
        str(output_path),
        pagesize=A4,
        leftMargin=1.4 * cm,
        rightMargin=1.4 * cm,
        topMargin=0.8 * cm,
        bottomMargin=0.8 * cm,
        title=profile["name"],
        author=profile["name"],
    )
    doc.build(story)


def collect_resume_text(profile: dict) -> str:
    parts = [profile["name"], profile["title"], profile["summary"]]
    parts.extend(skill["display"] for skill in profile["skills"])
    for role in profile["experience"]:
        parts.extend([role["role"], role["company"], role["dates"], role["location"]])
        parts.extend(role["bullets"])
    parts.extend(text for text, _ in profile["certifications"])
    parts.extend([profile["education"]["school"], profile["education"]["dates"]])
    return "\n".join(parts)


def compute_scores(profile: dict, jd_terms: set[str]) -> dict:
    resume_text = collect_resume_text(profile)
    matched = sorted(term for term in jd_terms if contains_term(term, resume_text))
    missing = sorted(term for term in jd_terms if term not in matched)
    match_score = (len(matched) / len(jd_terms) * 100.0) if jd_terms else 0.0

    structural_bonus = 15.0 if profile.get("summary") and profile.get("skills") and profile.get("experience") and profile.get("certifications") and profile.get("education") else 0.0
    ats_score = min(100.0, round((match_score * 0.85) + structural_bonus, 1))

    return {
        "match_score": round(match_score, 1),
        "ats_score": ats_score,
        "matched": matched,
        "missing": missing,
    }


def load_jd_text(args: argparse.Namespace) -> str:
    if args.jd_text:
        return args.jd_text
    if args.jd_file:
        return Path(args.jd_file).read_text(encoding="utf-8")
    if not sys.stdin.isatty():
        data = sys.stdin.read().strip()
        if data:
            return data
    if DEFAULT_JD_FILE.exists():
        return DEFAULT_JD_FILE.read_text(encoding="utf-8")
    raise SystemExit(
        "No job description provided. Use --jd-text, --jd-file, pipe JD text via stdin, or create jd_input.txt."
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Tailor the resume to a JD and generate a latest DOCX.")
    parser.add_argument("--jd-file", help="Path to a text file containing the job description.")
    parser.add_argument("--jd-text", help="Job description text passed directly on the command line.")
    parser.add_argument("--output", default=str(DEFAULT_OUTPUT), help="Output DOCX path.")
    parser.add_argument("--pdf-output", default=str(DEFAULT_PDF_OUTPUT), help="Output PDF path.")
    parser.add_argument(
        "--print-missing-only",
        action="store_true",
        help="Print only missing JD keywords instead of the full match report.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    jd_text = load_jd_text(args)
    jd_terms = extract_jd_terms(jd_text)
    tailored_profile = build_tailored_resume(jd_terms)
    scores = compute_scores(tailored_profile, jd_terms)

    output_path = Path(args.output).expanduser().resolve()
    pdf_output_path = Path(args.pdf_output).expanduser().resolve()
    build_document(tailored_profile, output_path)
    build_pdf_document(tailored_profile, pdf_output_path)

    print(f"Saved: {output_path.name}")
    print(f"Saved: {pdf_output_path.name}")
    print(f"Match score: {scores['match_score']:.1f}%")
    print(f"ATS heuristic score: {scores['ats_score']:.1f}%")
    print("Heuristic note: score is based on keyword overlap plus a small structural bonus.")
    if args.print_missing_only:
        print("Missing JD keywords:")
        for term in scores["missing"]:
            print(f"- {display_term(term)}")
        return

    print("Matched JD keywords:")
    for term in scores["matched"]:
        print(f"- {display_term(term)}")
    print("Missing JD keywords:")
    for term in scores["missing"]:
        print(f"- {display_term(term)}")


if __name__ == "__main__":
    main()
