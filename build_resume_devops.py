"""Generate a clean, single-page DevOps/SRE-focused ATS resume DOCX for Sai Charan Tumpuri."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = "Sai_Charan_Tumpuri_Resume_devops.docx"

doc = Document()

# Tight margins for single page
for section in doc.sections:
    section.top_margin = Cm(0.8)
    section.bottom_margin = Cm(0.8)
    section.left_margin = Cm(1.4)
    section.right_margin = Cm(1.4)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10)
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.space_before = Pt(0)


def add_para(text="", *, bold=False, italic=False, size=10, align=None,
             space_after=0, space_before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    return p


def add_horizontal_line(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "666666")
    pbdr.append(bottom)
    p_pr.append(pbdr)


def add_section_heading(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)
    add_horizontal_line(p)
    return p


def add_skill_line(label, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    r1 = p.add_run(f"{label}: ")
    r1.bold = True
    r1.font.size = Pt(10)
    r2 = p.add_run(value)
    r2.font.size = Pt(10)


def add_bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Cm(0.5)
    run = p.runs[0] if p.runs else p.add_run()
    if not p.runs:
        run = p.add_run(text)
    else:
        run.text = text
    run.font.size = Pt(10)


def _add_hyperlink(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1155CC")
    rPr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "20")
    rPr.append(sz)
    new_run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_bullet_with_link(text, url, link_text="Badge"):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Cm(0.5)
    r = p.add_run(text + "  |  ")
    r.font.size = Pt(10)
    _add_hyperlink(p, url, link_text)


def add_role_header(role, company, dates, location):
    # Role + dates row
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(17.5), WD_ALIGN_PARAGRAPH.RIGHT)
    r = p.add_run(f"{role} — {company}")
    r.bold = True
    r.font.size = Pt(10.5)
    r2 = p.add_run(f"\t{dates}")
    r2.bold = True
    r2.font.size = Pt(10)
    # Location line
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(2)
    r3 = p2.add_run(location)
    r3.italic = True
    r3.font.size = Pt(9.5)


# ---------- Header ----------
name = doc.add_paragraph()
name.alignment = WD_ALIGN_PARAGRAPH.CENTER
name.paragraph_format.space_after = Pt(0)
nr = name.add_run("SAI CHARAN TUMPURI")
nr.bold = True
nr.font.size = Pt(18)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(0)
tr = title.add_run("DevOps / Site Reliability Engineer")
tr.font.size = Pt(11)

contact = doc.add_paragraph()
contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
contact.paragraph_format.space_after = Pt(2)
cr = contact.add_run(
    "post2saicharan@gmail.com  |  +91 9390144066  |  Chennai, India  |  "
)
cr.font.size = Pt(9.5)
_add_hyperlink(contact, "https://github.com/saicharancodes", "GitHub")
cr2 = contact.add_run("  |  ")
cr2.font.size = Pt(9.5)
_add_hyperlink(contact, "https://www.linkedin.com/in/sai-charan-57b049232/", "LinkedIn")
cr3 = contact.add_run("  |  ")
cr3.font.size = Pt(9.5)
_add_hyperlink(contact, "https://saicharancodes.github.io/portfolio/", "Portfolio")

# ---------- Summary ----------
add_section_heading("Summary")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(2)
p.add_run(
    "Cloud DevOps/SRE engineer with 3+ years automating and scaling cloud "
    "infrastructure on GCP and AWS for 20+ teams at Sky (Comcast). Hands-on with "
    "Kubernetes (GKE/EKS), Terraform IaC at scale (300+ VMs), Linux and networking, "
    "and CI/CD across Jenkins, GitHub Actions, Cloud Build, and ArgoCD. Handle "
    "on-call incident response and SLO-driven observability. Built \u201cterrakit,\u201d an "
    "internal IaC self-service platform that cut infra ticket resolution time by 60%. "
    "AWS Solutions Architect \u2013 Associate and GCP Professional Cloud Architect certified."
).font.size = Pt(9.5)

# ---------- Skills (cleaned) ----------
add_section_heading("Technical Skills")
add_skill_line("Cloud", "GCP, AWS")
add_skill_line("Systems & Networking", "Linux, TCP/IP networking, DNS, load balancing, VPC")
add_skill_line("Containers & Orchestration", "Kubernetes (GKE, EKS), Docker, Helm")
add_skill_line("Infrastructure as Code", "Terraform, Ansible, Packer")
add_skill_line("CI/CD & GitOps", "Jenkins, GitHub Actions, Cloud Build, ArgoCD")
add_skill_line(
    "Observability & Incident Response",
    "Prometheus, Grafana, Cloud Monitoring, OpenTelemetry, distributed tracing, "
    "logging, on-call, SLO/SLA, postmortems",
)
add_skill_line("Languages", "Python, Go, Bash, Groovy, SQL")
add_skill_line("Data Platform", "BigQuery, Dataflow(Apache Beam), Airflow/Composer")

# ---------- Experience ----------
add_section_heading("Professional Experience")

add_role_header("DevOps Engineer II", "Comcast", "Mar 2024 – Present", "Chennai, India")
for b in [
    "Designed and operated Kubernetes-native CI/CD on GKE with ephemeral pod agents (Groovy + Python DSL) across 20+ pipelines, cutting idle compute by 30% and integrating E2E tests, artifact promotion, and vulnerability scans.",
    "Built \u201cterrakit,\u201d an internal Terraform abstraction with project-isolated remote state that lets engineers self-serve GCP infra (BigQuery, GKE, Dataflow); reduced infra ticket resolution time by 60% and unblocked 20+ data pipelines.",
    "Migrated legacy Dataflow and batch workloads to GKE using Helm and HPA, standardized chart templates, and drove capacity planning and scalability improvements by enforcing resource requests/limits — reducing job runtime by 25%.",
    "Owned OS and security lifecycle: rebuilt GCP golden images with Packer, led migration of 300+ VMs from CentOS 7 to CentOS 9 with zero SLA breaches, and decommissioned 40+ underutilized VMs to cut cost and CVE exposure.",
]:
    add_bullet(b)

add_role_header("DevOps Engineer I", "Comcast", "Jun 2023 – Mar 2024", "Chennai, India")
for b in [
    "Provisioned secure GCP infrastructure (VMs, IAM, BigQuery, GCS, VPC networking) using Terraform and Ansible with zero-touch Cloud Build CI/CD, cutting deployment time by 60%.",
    "Integrated InfraCost into GitHub PR checks for automated cost visibility on IaC changes, enabling proactive budget forecasting.",
    "Rebuilt 30+ legacy projects and decommissioned obsolete VMs; authored runbooks to support incident response and offshore knowledge transfer.",
    "Owned incident management for 300+ production Linux VMs on on-call rotation, monitoring CPU, memory, and 5xx alerts via Cloud Monitoring and Grafana; resolved P1\u2013P4 incidents against SLOs and authored postmortems to drive toil reduction.",
]:
    add_bullet(b)

# ---------- Certifications ----------
add_section_heading("Certifications")
add_bullet_with_link(
    "AWS Certified Solutions Architect \u2013 Associate (Sep 2025)",
    "https://www.credly.com/badges/1ef62e52-3d2c-4052-9254-19ea8275f0c1/public_url",
)
add_bullet_with_link(
    "Google Cloud Professional Cloud Architect (Feb 2026)",
    "https://www.credly.com/badges/0da7271f-b69b-4cfb-bc1c-9978cae5820c/public_url",
)

# ---------- Education ----------
add_section_heading("Education")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(0)
r = p.add_run("SASTRA University — B.Tech, Electronics & Communication Engineering")
r.bold = True
r.font.size = Pt(10)
p2 = doc.add_paragraph()
p2.paragraph_format.space_after = Pt(0)
r2 = p2.add_run("Jul 2019 – Jun 2023  |  GPA: 7.93")
r2.font.size = Pt(9.5)
r2.italic = True

doc.save(OUT)
print(f"Saved: {OUT}")
