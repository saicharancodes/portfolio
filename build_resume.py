"""Generate the single-page DevOps/SRE ATS resume (DOCX + PDF) for Sai Charan Tumpuri."""
import html

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm as RL_CM
from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate

OUT = "Sai_Charan_Tumpuri_Resume.docx"
PDF_OUT = "Sai_Charan_Tumpuri_Resume.pdf"

NAME = "SAI CHARAN TUMPURI"
TITLE = "DevOps / Site Reliability Engineer"
CONTACT_PREFIX = "post2saicharan@gmail.com  |  +91 9390144066  |  Chennai, India  |  "
CONTACT_LINKS = [
    ("GitHub", "https://github.com/saicharancodes"),
    ("LinkedIn", "https://www.linkedin.com/in/sai-charan-57b049232/"),
    ("Portfolio", "https://saicharancodes.github.io/portfolio/"),
]

SUMMARY = (
    "DevOps/SRE engineer with 3+ years of automation and scalable cloud infrastructure "
    "on GCP and AWS, supporting 20+ teams at Comcast (Sky). Hands-on with Kubernetes "
    "(GKE/EKS), Terraform infrastructure as code at scale (300+ VMs), Linux, networking, "
    "and CI/CD across Jenkins, GitHub Actions, Cloud Build, and ArgoCD. Experienced in "
    "on-call incident response, observability, monitoring, and SLO-driven reliability. "
    "Built TerraKit, an internal IaC self-service platform that cut infra ticket "
    "resolution time by 60%. AWS Certified Solutions Architect \u2013 Associate and "
    "Google Cloud Professional Cloud Architect."
)

SKILLS = [
    ("Cloud", "GCP, AWS"),
    ("Containers & Orchestration", "Kubernetes (GKE, EKS), Docker, Helm"),
    ("Infrastructure as Code", "Terraform, Ansible, Packer"),
    ("CI/CD & GitOps", "Jenkins, GitHub Actions, Cloud Build, ArgoCD"),
    ("Observability & Monitoring", "Prometheus, Grafana, Cloud Monitoring, OpenTelemetry"),
    ("Systems & Networking", "Linux, TCP/IP, DNS, VPC, Load Balancers"),
    ("Languages", "Python, Go, Bash/Shell, Groovy, SQL"),
    ("Data Platform", "BigQuery, Dataflow (Apache Beam), Airflow/Composer"),
    ("AI Tools", "GitHub Copilot, Claude Code, Antigravity"),
]

EXPERIENCE = [
    {
        "role": "DevOps Engineer II",
        "company": "Comcast",
        "dates": "Mar 2024 \u2013 Present",
        "location": "Chennai, India",
        "bullets": [
            "Designed and operated **Kubernetes**-native **CI/CD** on **GKE** with ephemeral pod agents (**Groovy** + **Python** DSL) across 20+ pipelines, automating E2E tests, artifact promotion, and vulnerability scans and cutting idle compute by 30% (**cost optimization**).",
            "Built **TerraKit**, an internal **Terraform** abstraction with project-isolated remote state that lets engineers self-serve **GCP** infrastructure (**BigQuery**, **GKE**, **Dataflow**); reduced infra ticket resolution time by 60% (**toil reduction**) and unblocked 20+ data pipelines.",
            "Designed and deployed a production-grade **FastAPI** RAG platform on **Kubernetes (GKE)** using Deployments, StatefulSets, Secrets, PVCs, readiness/liveness probes, and resource limits, backed by **PostgreSQL/pgvector**, **Prometheus** metrics, and **Docker** CI/CD publishing multi-architecture images to **GitHub Container Registry**.",
            "Owned OS and security lifecycle: rebuilt GCP golden images with **Packer**, led migration of 300+ **Linux** VMs from CentOS 7 to CentOS 9 with zero SLA breaches, and decommissioned 40+ underutilized VMs to cut cost and CVE exposure.",
        ],
    },
    {
        "role": "DevOps Engineer I",
        "company": "Comcast",
        "dates": "Jun 2023 \u2013 Mar 2024",
        "location": "Chennai, India",
        "bullets": [
            "Provisioned secure **GCP** infrastructure (VMs, IAM, **BigQuery**, GCS, **VPC networking**) using **Terraform** and **Ansible** with zero-touch **Cloud Build** CI/CD automation, cutting deployment time by 60%.",
            "Integrated **InfraCost** into **GitHub** PR checks for automated cost visibility on **infrastructure as code** changes, enabling proactive budget forecasting and **cost optimization**.",
            "Rebuilt 30+ legacy projects and decommissioned obsolete VMs; authored runbooks to support **incident response** and offshore knowledge transfer.",
            "Owned **incident response** for 300+ production **Linux** VMs on **on-call** rotation, monitoring CPU, memory, and 5xx alerts via **Cloud Monitoring** and **Grafana** (**observability**, logging); resolved P1\u2013P4 incidents against **SLOs** and authored **postmortems** to drive toil reduction.",
        ],
    },
]

CERTIFICATIONS = [
    (
        "AWS Certified Solutions Architect \u2013 Associate (Sep 2025)",
        "https://www.credly.com/badges/1ef62e52-3d2c-4052-9254-19ea8275f0c1/public_url",
    ),
    (
        "Google Cloud Professional Cloud Architect (Feb 2026)",
        "https://www.credly.com/badges/0da7271f-b69b-4cfb-bc1c-9978cae5820c/public_url",
    ),
]

EDUCATION_SCHOOL = "SASTRA University \u2014 B.Tech, Electronics & Communication Engineering"
EDUCATION_DATES = "Jul 2019 \u2013 Jun 2023  |  GPA: 7.93"


def split_bold(text):
    """Split text on **...** markers into (segment, is_bold) pairs."""
    return [(part, index % 2 == 1) for index, part in enumerate(text.split("**")) if part]

doc = Document()

# Tight margins for single page
for section in doc.sections:
    section.top_margin = Cm(0.95)
    section.bottom_margin = Cm(0.95)
    section.left_margin = Cm(1.4)
    section.right_margin = Cm(1.4)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10)
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.space_before = Pt(0)


def add_para(text="", *, bold=False, italic=False, size=10, align=None,
             space_after=0, space_before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    return p


def add_horizontal_line(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "666666")
    pbdr.append(bottom)
    p_pr.append(pbdr)


def add_section_heading(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)
    add_horizontal_line(p)
    return p


def add_skill_line(label, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r1 = p.add_run(f"{label}: ")
    r1.bold = True
    r1.font.size = Pt(10)
    r2 = p.add_run(value)
    r2.font.size = Pt(10)


def add_bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Cm(0.5)
    for segment, bold in split_bold(text):
        run = p.add_run(segment)
        run.bold = bold
        run.font.size = Pt(10)


def _add_hyperlink(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1155CC")
    rPr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "20")
    rPr.append(sz)
    new_run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_bullet_with_link(text, url, link_text="Badge"):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Cm(0.5)
    r = p.add_run(text + "  |  ")
    r.font.size = Pt(10)
    _add_hyperlink(p, url, link_text)


def add_role_header(role, company, dates, location):
    # Role + dates row
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(17.5), WD_ALIGN_PARAGRAPH.RIGHT)
    r = p.add_run(f"{role} — {company}")
    r.bold = True
    r.font.size = Pt(10.5)
    r2 = p.add_run(f"\t{dates}")
    r2.bold = True
    r2.font.size = Pt(10)
    # Location line
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(3)
    r3 = p2.add_run(location)
    r3.italic = True
    r3.font.size = Pt(9.5)


# ---------- Header ----------
name = doc.add_paragraph()
name.alignment = WD_ALIGN_PARAGRAPH.CENTER
name.paragraph_format.space_after = Pt(0)
nr = name.add_run(NAME)
nr.bold = True
nr.font.size = Pt(18)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(0)
tr = title.add_run(TITLE)
tr.font.size = Pt(11)

contact = doc.add_paragraph()
contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
contact.paragraph_format.space_after = Pt(2)
cr = contact.add_run(CONTACT_PREFIX)
cr.font.size = Pt(9.5)
for index, (label, url) in enumerate(CONTACT_LINKS):
    _add_hyperlink(contact, url, label)
    if index < len(CONTACT_LINKS) - 1:
        spacer = contact.add_run("  |  ")
        spacer.font.size = Pt(9.5)

# ---------- Summary ----------
add_section_heading("Summary")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(2)
p.add_run(SUMMARY).font.size = Pt(9.5)

# ---------- Skills (cleaned) ----------
add_section_heading("Technical Skills")
for label, value in SKILLS:
    add_skill_line(label, value)

# ---------- Experience ----------
add_section_heading("Professional Experience")

for role in EXPERIENCE:
    add_role_header(role["role"], role["company"], role["dates"], role["location"])
    for b in role["bullets"]:
        add_bullet(b)

# ---------- Certifications ----------
add_section_heading("Certifications")
for cert_text, cert_url in CERTIFICATIONS:
    add_bullet_with_link(cert_text, cert_url)

# ---------- Education ----------
add_section_heading("Education")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(0)
r = p.add_run(EDUCATION_SCHOOL)
r.bold = True
r.font.size = Pt(10)
p2 = doc.add_paragraph()
p2.paragraph_format.space_after = Pt(0)
r2 = p2.add_run(EDUCATION_DATES)
r2.font.size = Pt(9.5)
r2.italic = True

doc.save(OUT)
print(f"Saved: {OUT}")


def build_pdf(path):
    base = getSampleStyleSheet()
    name_style = ParagraphStyle("Name", parent=base["Title"], fontName="Helvetica-Bold",
                                fontSize=18, leading=20, alignment=TA_CENTER, spaceAfter=2)
    title_style = ParagraphStyle("Subtitle", parent=base["Normal"], fontName="Helvetica",
                                 fontSize=11, leading=13, alignment=TA_CENTER, spaceAfter=2)
    contact_style = ParagraphStyle("Contact", parent=base["Normal"], fontName="Helvetica",
                                   fontSize=9, leading=11, alignment=TA_CENTER, spaceAfter=5)
    section_style = ParagraphStyle("Section", parent=base["Heading2"], fontName="Helvetica-Bold",
                                   fontSize=11, leading=13, textColor="#1F2A44",
                                   spaceBefore=5, spaceAfter=2)
    body_style = ParagraphStyle("Body", parent=base["Normal"], fontName="Helvetica",
                                fontSize=9.2, leading=11.4, spaceAfter=2)
    role_style = ParagraphStyle("Role", parent=base["Normal"], fontName="Helvetica",
                                fontSize=10.5, leading=12.5, spaceBefore=5, spaceAfter=1)
    location_style = ParagraphStyle("Location", parent=base["Normal"], fontName="Helvetica-Oblique",
                                                                        fontSize=9.5, leading=11.5, spaceAfter=3)
    bullet_style = ParagraphStyle("Bullet", parent=body_style, leftIndent=12,
                                                                      bulletIndent=2, spaceAfter=2)

    def link(text, url):
        return f'<font color="#1155CC"><link href="{html.escape(url)}"><u>{html.escape(text)}</u></link></font>'

    story = [Paragraph(html.escape(NAME), name_style),
             Paragraph(html.escape(TITLE), title_style)]

    contact = html.escape(CONTACT_PREFIX) + "  |  ".join(link(l, u) for l, u in CONTACT_LINKS)
    story.append(Paragraph(contact, contact_style))

    def section(heading):
        story.append(Paragraph(heading.upper(), section_style))
        story.append(HRFlowable(width="100%", thickness=0.8, color="#666666",
                                spaceBefore=0, spaceAfter=3))

    section("Summary")
    story.append(Paragraph(html.escape(SUMMARY), body_style))

    section("Technical Skills")
    for label, value in SKILLS:
        story.append(Paragraph(f"<b>{html.escape(label)}:</b> {html.escape(value)}", body_style))

    section("Professional Experience")
    for role in EXPERIENCE:
        header = (f"<b>{html.escape(role['role'])}</b>  |  "
                  f"<b>{html.escape(role['company'])}</b>  |  "
                  f"{html.escape(role['dates'])}")
        story.append(Paragraph(header, role_style))
        story.append(Paragraph(html.escape(role["location"]), location_style))
        for b in role["bullets"]:
            rich = "".join(f"<b>{html.escape(s)}</b>" if bold else html.escape(s)
                           for s, bold in split_bold(b))
            story.append(Paragraph(rich, bullet_style, bulletText="\u2022"))

    section("Certifications")
    for cert_text, cert_url in CERTIFICATIONS:
        story.append(Paragraph(f"{html.escape(cert_text)}  |  {link('Badge', cert_url)}",
                               bullet_style, bulletText="\u2022"))

    section("Education")
    story.append(Paragraph(f"<b>{html.escape(EDUCATION_SCHOOL)}</b>", body_style))
    story.append(Paragraph(f"<i>{html.escape(EDUCATION_DATES)}</i>", body_style))

    SimpleDocTemplate(path, pagesize=A4,
                      leftMargin=1.4 * RL_CM, rightMargin=1.4 * RL_CM,
                      topMargin=0.95 * RL_CM, bottomMargin=0.95 * RL_CM,
                      title=NAME, author=NAME).build(story)


build_pdf(PDF_OUT)
print(f"Saved: {PDF_OUT}")
